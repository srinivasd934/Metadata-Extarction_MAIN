from pathlib import Path
import pandas as pd
import os
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe\tesseract.exe'
from PIL import Image
from docx import Document
from pypdf import PdfReader
from .config import SUPPORTED_EXTENSIONS


def extract_docx(path: Path) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                tables.append(" | ".join(cells))
    return "\n".join(paragraphs + tables).strip()


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def extract_image(path: Path) -> str:
    return pytesseract.image_to_string(Image.open(path)).strip()


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore").strip()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return extract_image(path)
    if suffix == ".txt":
        return extract_txt(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def normalize_filename(name: str) -> str:
    value = str(name).strip()
    for suffix in [".docx", ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".txt"]:
        if value.lower().endswith(suffix):
            value = value[:-len(suffix)]
    return value.strip()


def build_document_dataframe(train_dir: Path, test_dir: Path) -> pd.DataFrame:
    rows = []
    for split, folder in [("train", train_dir), ("test", test_dir)]:
        if not folder.exists():
            continue
        for path in sorted(folder.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            try:
                text = extract_text(path)
                status = "success" if text.strip() else "empty"
            except Exception as exc:
                text = ""
                status = f"error: {exc}"
            rows.append({
                "File Name": path.name,
                "File Stem": normalize_filename(path.name),
                "Split": split,
                "Text": text,
                "Character Count": len(text),
                "Status": status,
            })
    return pd.DataFrame(rows)


